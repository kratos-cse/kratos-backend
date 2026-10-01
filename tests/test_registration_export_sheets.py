"""Tests for de-congested team member columns and multi-sheet registration exports."""
import io
import uuid
import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.models.enums import (
    MemberRegistrationMode,
    PaymentStatus,
    PaymentType,
    RegistrationMode,
    RegistrationStatus,
    TeamMemberRole,
    TeamMemberStatus,
    TeamStatus,
)
from app.models.event import Event, EventRegistrationRule
from app.models.event_content import EventRegistrationField, RegistrationFieldResponse
from app.models.enums import RegistrationFieldScope, RegistrationFieldSource, RegistrationFieldType
from app.models.payment import Payment
from app.models.profile import Profile
from app.models.registration import Registration
from app.models.team import Team, TeamMember
from app.models.admin import AdminUser, Role
from app.models.event_admin_assignment import EventAdminAssignment
from app.services import admin_ops_service as ops
from tests.conftest import _make_event, _make_user_profile, requires_db


@requires_db
@pytest.mark.asyncio
async def test_single_event_export_predetermined_team_size(db):
    """Event with predetermined team size 3 should produce Member 1, Member 2, Member 3 columns."""
    event = await _make_event(db, name=f"Trio Challenge {uuid.uuid4().hex[:6]}")
    rules_res = await db.execute(select(EventRegistrationRule).where(EventRegistrationRule.event_id == event.id))
    rules = rules_res.scalar_one()
    rules.registration_mode = RegistrationMode.TEAM_ONLY
    rules.team_min_size = 3
    rules.team_max_size = 3
    rules.required_member_count = 3
    rules.allow_individual = False
    await db.flush()

    leader = await _make_user_profile(db, email=f"lead-{uuid.uuid4().hex}@test.edu")
    m2 = await _make_user_profile(db, email=f"m2-{uuid.uuid4().hex}@test.edu")

    team = Team(
        event_id=event.id,
        name="Alpha Squad",
        leader_profile_id=leader.id,
        status=TeamStatus.PAID,
    )
    db.add(team)
    await db.flush()

    tm1 = TeamMember(
        team_id=team.id,
        event_id=event.id,
        profile_id=leader.id,
        role=TeamMemberRole.LEADER,
        status=TeamMemberStatus.ACTIVE,
    )
    tm2 = TeamMember(
        team_id=team.id,
        event_id=event.id,
        profile_id=m2.id,
        role=TeamMemberRole.MEMBER,
        status=TeamMemberStatus.ACTIVE,
    )
    # Leader-entered 3rd member
    tm3 = TeamMember(
        team_id=team.id,
        event_id=event.id,
        profile_id=None,
        full_name="Third Player",
        phone="9876543210",
        role=TeamMemberRole.MEMBER,
        status=TeamMemberStatus.ACTIVE,
    )
    db.add_all([tm1, tm2, tm3])

    reg = Registration(
        event_id=event.id,
        team_id=team.id,
        status=RegistrationStatus.CONFIRMED,
    )
    db.add(reg)
    await db.commit()

    headers, rows = await ops._get_registration_export_data(db, event_id=event.id)

    assert "Member 1" in headers
    assert "Member 2" in headers
    assert "Member 3" in headers
    assert "Member 4" not in headers

    assert len(rows) == 1
    row = rows[0]
    idx_m1 = headers.index("Member 1")
    idx_m2 = headers.index("Member 2")
    idx_m3 = headers.index("Member 3")

    # Member 1 is leader
    assert leader.full_name in row[idx_m1]
    assert "LEADER" in row[idx_m1]

    # Member 2 is m2
    assert m2.full_name in row[idx_m2]
    assert "MEMBER" in row[idx_m2]

    # Member 3 is leader-entered
    assert "Third Player" in row[idx_m3]
    assert "9876543210" in row[idx_m3]


@requires_db
@pytest.mark.asyncio
async def test_multi_sheet_workbook_custom_columns_per_event(db):
    """Multi-event XLSX export has per-event sheet with tailored member columns."""
    # Event 1: Duos (2 members)
    ev_duo = await _make_event(db, name=f"Duo Cup {uuid.uuid4().hex[:6]}")
    r1_res = await db.execute(select(EventRegistrationRule).where(EventRegistrationRule.event_id == ev_duo.id))
    r1 = r1_res.scalar_one()
    r1.team_min_size = 2
    r1.team_max_size = 2
    r1.required_member_count = 2

    # Event 2: Solo (1 member)
    ev_solo = await _make_event(db, name=f"Solo Coding {uuid.uuid4().hex[:6]}")
    r2_res = await db.execute(select(EventRegistrationRule).where(EventRegistrationRule.event_id == ev_solo.id))
    r2 = r2_res.scalar_one()
    r2.team_min_size = 1
    r2.team_max_size = 1
    r2.registration_mode = RegistrationMode.INDIVIDUAL_ONLY

    await db.commit()

    buf = await ops.export_registrations_xlsx(db)
    wb = load_workbook(io.BytesIO(buf.getvalue()))

    # Find sheets
    sheet_names = wb.sheetnames
    safe_duo = ev_duo.name[:31]
    safe_solo = ev_solo.name[:31]

    assert safe_duo in sheet_names
    assert safe_solo in sheet_names

    ws_duo = wb[safe_duo]
    duo_headers = [cell.value for cell in ws_duo[1]]
    assert "Member 1" in duo_headers
    assert "Member 2" in duo_headers
    assert "Member 3" not in duo_headers

    ws_solo = wb[safe_solo]
    solo_headers = [cell.value for cell in ws_solo[1]]
    assert "Member 1" not in solo_headers


def test_unit_member_columns_count_predetermined():
    """Unit test: predetermined team size returns exact member count."""
    event = Event(id=uuid.uuid4(), name="Trio Code")
    rules = EventRegistrationRule(event_id=event.id, team_min_size=3, team_max_size=3)
    event.rules = rules

    count = ops._get_event_member_columns_count(event, [])
    assert count == 3

    # Solo event
    solo_ev = Event(id=uuid.uuid4(), name="Solo")
    solo_rules = EventRegistrationRule(event_id=solo_ev.id, team_min_size=1, team_max_size=1)
    solo_ev.rules = solo_rules
    assert ops._get_event_member_columns_count(solo_ev, []) == 0


def test_unit_sorted_active_team_members():
    """Unit test: leader is first, left/removed members are excluded."""
    team = Team(id=uuid.uuid4(), name="Bravo")
    m_member = TeamMember(id=uuid.uuid4(), role=TeamMemberRole.MEMBER, status=TeamMemberStatus.ACTIVE)
    m_leader = TeamMember(id=uuid.uuid4(), role=TeamMemberRole.LEADER, status=TeamMemberStatus.ACTIVE)
    m_removed = TeamMember(id=uuid.uuid4(), role=TeamMemberRole.MEMBER, status=TeamMemberStatus.REMOVED)
    m_left = TeamMember(id=uuid.uuid4(), role=TeamMemberRole.MEMBER, status=TeamMemberStatus.LEFT)

    team.members = [m_member, m_removed, m_leader, m_left]
    active = ops._sorted_active_team_members(team)

    assert len(active) == 2
    assert active[0].role == TeamMemberRole.LEADER
    assert active[1].role == TeamMemberRole.MEMBER


def test_unit_multi_sheet_workbook_bytes_custom_headers():
    """Unit test: _multi_sheet_workbook_bytes generates distinct headers for distinct sheets."""
    sheets_data = {
        "Duo Event": (["Reg ID", "Member 1", "Member 2"], [["1", "Alice", "Bob"]]),
        "Trio Event": (["Reg ID", "Member 1", "Member 2", "Member 3"], [["2", "Charlie", "Dave", "Eve"]]),
    }
    buf = ops._multi_sheet_workbook_bytes(sheets_data=sheets_data)
    wb = load_workbook(io.BytesIO(buf.getvalue()))

    assert "Duo Event" in wb.sheetnames
    assert "Trio Event" in wb.sheetnames

    duo_h = [cell.value for cell in wb["Duo Event"][1]]
    assert duo_h == ["Reg ID", "Member 1", "Member 2"]

    trio_h = [cell.value for cell in wb["Trio Event"][1]]
    assert trio_h == ["Reg ID", "Member 1", "Member 2", "Member 3"]
